"""
LaTeX 数学公式转换器
"""

from io import BytesIO
from typing import Optional

import requests
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from urllib.parse import quote

from ..equation_labels import EquationRegistry, strip_label
from ..security import MAX_IMAGE_BYTES, is_allowed_codecogs_url
from .base import ElementConverter
from .mermaid import _is_valid_image_payload

CODECOGS_PNG_BASE = "https://latex.codecogs.com/png.latex?"
MAX_LATEX_SOURCE_LEN = 4096


def build_codecogs_url(latex: str, inline: bool = False) -> str:
    """构造 CodeCogs PNG 渲染 URL"""
    source = (latex or "").strip()
    if inline:
        source = f"\\inline {source}"
    encoded = quote(source, safe="")
    return f"{CODECOGS_PNG_BASE}{encoded}"


class MathConverter(ElementConverter):
    """将 LaTeX 公式渲染为 PNG 并嵌入 DOCX"""

    def convert(self, token) -> None:
        """块级 math_block token"""
        if not self.document:
            raise ValueError("Document not set")
        latex = token.content if hasattr(token, "content") else ""
        if not latex.strip():
            return

        registry = self._equation_registry()
        equation_number: Optional[int] = None
        if registry is not None:
            latex, label_id = strip_label(latex)
            if not latex.strip():
                return
            equation_number = registry.next_number()
            if label_id:
                registry.register(label_id, equation_number)

        if self._try_embed(latex, block=True, equation_number=equation_number):
            return
        self._fallback_block(latex, equation_number=equation_number)

    def convert_in_paragraph(self, paragraph: Paragraph, token) -> None:
        """行内 math_inline token"""
        latex = token.content if hasattr(token, "content") else ""
        if not latex.strip():
            return
        if self._try_embed_inline(paragraph, latex):
            return
        run = paragraph.add_run(f"${latex}$")
        run.italic = True
        run.font.name = "Consolas"
        run.font.size = Pt(10)

    def _equation_registry(self) -> Optional[EquationRegistry]:
        if self.base_converter and hasattr(self.base_converter, "_equation_registry"):
            return self.base_converter._equation_registry
        return None

    def _try_embed(
        self, latex: str, block: bool = False, equation_number: Optional[int] = None
    ) -> bool:
        image_data = self._fetch_formula_image(latex, inline=not block)
        if not image_data:
            return False
        try:
            if block:
                self._embed_block_image(image_data, equation_number)
            else:
                paragraph = self.doc.add_paragraph()
                self._embed_inline_image(paragraph, image_data)
            return True
        except Exception:
            return False

    def _try_embed_inline(self, paragraph: Paragraph, latex: str) -> bool:
        image_data = self._fetch_formula_image(latex, inline=True)
        if not image_data:
            return False
        try:
            self._embed_inline_image(paragraph, image_data)
            return True
        except Exception:
            return False

    def _fetch_formula_image(self, latex: str, inline: bool = False) -> Optional[bytes]:
        source = (latex or "").strip()
        if not source or len(source) > MAX_LATEX_SOURCE_LEN:
            return None
        url = build_codecogs_url(source, inline=inline)
        if not is_allowed_codecogs_url(url):
            return None
        try:
            response = requests.get(url, timeout=15, stream=True)
            if response.status_code != 200:
                return None
            content_type = response.headers.get("Content-Type", "")
            if content_type and "image" not in content_type:
                return None
            chunks = []
            total = 0
            for chunk in response.iter_content(chunk_size=8192):
                total += len(chunk)
                if total > MAX_IMAGE_BYTES:
                    return None
                chunks.append(chunk)
            data = b"".join(chunks)
            return data if _is_valid_image_payload(data) else None
        except Exception:
            return None

    def _embed_block_image(
        self, image_data: bytes, equation_number: Optional[int] = None
    ) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(BytesIO(image_data), width=Inches(4.0))

        if equation_number is not None:
            caption = self.doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_run = caption.add_run(f"({equation_number})")
            cap_run.italic = True
            cap_run.font.size = Pt(9)
            cap_run.font.color.rgb = RGBColor(102, 102, 102)

    def _embed_inline_image(self, paragraph: Paragraph, image_data: bytes) -> None:
        run = paragraph.add_run()
        run.add_picture(BytesIO(image_data), height=Inches(0.22))

    def _fallback_block(
        self, latex: str, equation_number: Optional[int] = None
    ) -> None:
        note_para = self.doc.add_paragraph()
        note_run = note_para.add_run("（公式渲染失败，已保留 LaTeX 源码）")
        note_run.italic = True
        note_run.font.size = Pt(9)
        note_run.font.color.rgb = RGBColor(128, 128, 128)

        code_para = self.doc.add_paragraph()
        try:
            code_para.style = "Code"
        except KeyError:
            pass
        code_para.add_run(latex)

        if equation_number is not None:
            caption = self.doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_run = caption.add_run(f"({equation_number})")
            cap_run.italic = True
            cap_run.font.size = Pt(9)
            cap_run.font.color.rgb = RGBColor(102, 102, 102)
