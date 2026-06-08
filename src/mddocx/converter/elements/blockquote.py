"""
引用块转换器模块，处理引用块的转换
"""

from typing import Any, Dict, Tuple

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

from .base import ElementConverter


class BlockquoteConverter(ElementConverter):
    """处理引用块的转换器"""

    def __init__(self, base_converter=None):
        super().__init__(base_converter)

    def convert(self, tokens: Tuple[Any, Any]) -> None:
        """转换引用块元素

        Args:
            tokens: (开始标记, 内容标记) 的元组
        """
        if not self.document:
            raise ValueError("Document not set")

        quote_token, content_token = tokens

        level = len(quote_token.markup) if hasattr(quote_token, "markup") else 1

        style_name = "Quote" if level == 1 else f"Quote{level}"
        self._ensure_quote_style(style_name, level)

        paragraph = self.document.add_paragraph()
        paragraph.style = self.document.styles[style_name]

        if not content_token:
            paragraph.add_run("")
            return

        link_converter = None
        if self.base_converter and "link" in self.base_converter.converters:
            link_converter = self.base_converter.converters.get("link")

        self._process_inline_children(
            paragraph, content_token.children, link_converter
        )

    def _process_inline_children(
        self, paragraph: Paragraph, children, link_converter
    ) -> None:
        """处理引用块内联内容（链接、删除线、行内代码等）"""
        current_text = ""
        current_style: Dict[str, bool] = {
            "bold": False,
            "italic": False,
            "strike": False,
        }

        i = 0
        while i < len(children):
            child = children[i]

            if child.type == "text":
                text = child.content.replace("\n", " ")
                if text.endswith(" "):
                    text = text[:-1]
                current_text += text
                i += 1
            elif child.type == "link_open":
                if current_text:
                    self._add_text_with_style(paragraph, current_text, current_style)
                    current_text = ""

                link_content = None
                j = i + 1
                while j < len(children) and children[j].type != "link_close":
                    if children[j].type == "text":
                        link_content = children[j]
                    j += 1

                if link_content and link_converter:
                    link_text = (
                        link_content.content
                        if hasattr(link_content, "content")
                        else None
                    )
                    link_converter.convert_in_paragraph(
                        paragraph, child, current_style.copy(), link_text
                    )
                elif link_content:
                    self._add_text_with_style(
                        paragraph, link_content.content, current_style
                    )

                i = j + 1 if j < len(children) else i + 1
            elif child.type == "link_close":
                i += 1
            elif child.type == "strong_open":
                if current_text:
                    self._add_text_with_style(paragraph, current_text, current_style)
                    current_text = ""
                current_style["bold"] = True
                i += 1
            elif child.type == "strong_close":
                if current_text:
                    self._add_text_with_style(paragraph, current_text, current_style)
                    current_text = ""
                current_style["bold"] = False
                i += 1
            elif child.type == "em_open":
                if current_text:
                    self._add_text_with_style(paragraph, current_text, current_style)
                    current_text = ""
                current_style["italic"] = True
                i += 1
            elif child.type == "em_close":
                if current_text:
                    self._add_text_with_style(paragraph, current_text, current_style)
                    current_text = ""
                current_style["italic"] = False
                i += 1
            elif child.type == "s_open":
                if current_text:
                    self._add_text_with_style(paragraph, current_text, current_style)
                    current_text = ""
                current_style["strike"] = True
                i += 1
            elif child.type == "s_close":
                if current_text:
                    self._add_text_with_style(paragraph, current_text, current_style)
                    current_text = ""
                current_style["strike"] = False
                i += 1
            elif child.type == "code_inline":
                if current_text:
                    self._add_text_with_style(paragraph, current_text, current_style)
                    current_text = ""
                self._add_inline_code(paragraph, child.content, current_style.copy())
                i += 1
            elif child.type == "softbreak":
                current_text += " "
                i += 1
            else:
                i += 1

        if current_text:
            self._add_text_with_style(paragraph, current_text, current_style)

    def _add_text_with_style(
        self, paragraph: Paragraph, text: str, style: Dict[str, bool]
    ) -> None:
        """添加带样式的文本"""
        run = paragraph.add_run(text)
        run.bold = style["bold"]
        run.italic = style["italic"]
        run.font.strike = style["strike"]

    def _add_inline_code(
        self, paragraph: Paragraph, code_text: str, style: Dict[str, bool]
    ) -> None:
        """添加行内代码"""
        run = paragraph.add_run(code_text)
        run.bold = style.get("bold", False)
        run.italic = style.get("italic", False)
        run.font.strike = style.get("strike", False)
        run.font.name = "Consolas"
        run.font.size = Pt(10)

    def _ensure_quote_style(self, style_name: str, level: int) -> None:
        """确保引用块样式存在"""
        if style_name not in self.document.styles:
            style = self.document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
            style.font.size = Pt(12)
            style.font.color.rgb = RGBColor(102, 102, 102)
            style.paragraph_format.left_indent = Pt(30 * level)
            style.paragraph_format.space_before = Pt(6)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
