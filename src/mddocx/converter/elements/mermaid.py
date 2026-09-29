"""
Mermaid 流程图转换器
"""

import base64
import json
from io import BytesIO
from typing import Any, Optional

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from ..render_cache import RenderCache
from ..security import MAX_IMAGE_BYTES, fetch_remote_bytes, is_allowed_mermaid_ink_url
from .base import ElementConverter

SUPPORTED_DIAGRAM_PREFIXES = (
    "graph ",
    "flowchart ",
    "sequencediagram",
    "gantt",
    "statediagram-v2",
    "statediagram",
    "classdiagram",
    "pie",
)


def mermaid_diagram_kind(source: str) -> str:
    """返回图表类型中文标签（用于 DOCX 说明文字）。"""
    if not source or not source.strip():
        return "图表"
    first = source.strip().split("\n", 1)[0].strip().lower()
    if first.startswith("sequencediagram"):
        return "时序图"
    if first.startswith("gantt"):
        return "甘特图"
    if first.startswith("statediagram"):
        return "状态图"
    if first.startswith("classdiagram"):
        return "类图"
    if first.startswith("pie"):
        return "饼图"
    if first.startswith(("graph ", "flowchart ")):
        return "流程图"
    return "图表"


_IMAGE_MAGIC_PREFIXES = (
    b"\x89PNG\r\n\x1a\n",  # PNG
    b"\xff\xd8\xff",  # JPEG
    b"GIF87a",
    b"GIF89a",  # GIF
    b"II*\x00",
    b"MM\x00*",  # TIFF
    b"BM",  # BMP
)


def _is_valid_image_payload(data: bytes) -> bool:
    """按魔数判断渲染结果是否为可嵌入的图片。

    非图片响应（HTML 错误页、任意二进制）必须在嵌入前被拒掉，否则
    ``add_picture`` 抛错后只能走回退路径，且缓存里会存下垃圾数据。
    """
    if not data:
        return False
    if any(data.startswith(prefix) for prefix in _IMAGE_MAGIC_PREFIXES):
        return True
    # WEBP: RIFF....WEBP
    return len(data) >= 12 and data.startswith(b"RIFF") and data[8:12] == b"WEBP"


def is_supported_mermaid_diagram(source: str) -> bool:
    """支持 graph / flowchart / sequence / gantt / state / class / pie"""
    if not source or not source.strip():
        return False
    first = source.strip().split("\n", 1)[0].strip().lower()
    return first.startswith(SUPPORTED_DIAGRAM_PREFIXES)


def build_mermaid_ink_url(diagram: str, image_type: str = "png") -> str:
    """构造 mermaid.ink 渲染 URL"""
    state = {
        "code": diagram.strip(),
        "mermaid": json.dumps({"theme": "default"}),
    }
    payload = json.dumps(state, separators=(",", ":"))
    encoded = base64.urlsafe_b64encode(payload.encode("utf-8")).decode("ascii")
    return f"https://mermaid.ink/img/{encoded}?type={image_type}"


class MermaidConverter(ElementConverter):
    """将 Mermaid 图表渲染为图片并嵌入 DOCX"""

    def __init__(self, base_converter: Optional[Any] = None) -> None:
        super().__init__(base_converter)
        # 同一文档内重复的图表不再重复请求 mermaid.ink
        self._render_cache = RenderCache()

    def convert(self, token: Any) -> None:
        if not self.document:
            raise ValueError("Document not set")

        source = token.content if hasattr(token, "content") else ""

        if not is_supported_mermaid_diagram(source):
            self._fallback_as_code(token, source, unsupported=True)
            return

        image_data = self._fetch_diagram_image(source)
        if image_data:
            try:
                self._embed_image(image_data, mermaid_diagram_kind(source))
                return
            except Exception:
                pass
        self._fallback_as_code(token, source)

    def _fetch_diagram_image(self, diagram: str) -> Optional[bytes]:
        cached = self._render_cache.get(diagram)
        if cached is not None:
            return cached

        url = build_mermaid_ink_url(diagram)
        data = fetch_remote_bytes(
            url,
            validator=is_allowed_mermaid_ink_url,
            timeout=15,
            max_bytes=MAX_IMAGE_BYTES,
            content_type_prefix="image",
        )
        if not data or not _is_valid_image_payload(data):
            return None
        self._render_cache.put(diagram, data)
        return data

    def _embed_image(self, image_data: bytes, kind: str = "流程图") -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(BytesIO(image_data), width=Inches(5.5))

        caption = self.doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = caption.add_run(f"Mermaid {kind}")
        cap_run.italic = True
        cap_run.font.size = Pt(9)
        cap_run.font.color.rgb = RGBColor(102, 102, 102)

    def _fallback_as_code(self, token: Any, source: str, unsupported: bool = False) -> None:
        note = "（不支持的 Mermaid 类型，已保留源码）" if unsupported else "（Mermaid 渲染失败，已保留源码）"
        note_para = self.doc.add_paragraph()
        note_run = note_para.add_run(note)
        note_run.italic = True
        note_run.font.size = Pt(9)
        note_run.font.color.rgb = RGBColor(128, 128, 128)

        code_converter = None
        if self.base_converter:
            code_converter = self.base_converter.converters.get("code")
        if code_converter:
            code_converter.set_document(self.doc)
            code_converter.convert(token)
        else:
            paragraph = self.doc.add_paragraph()
            try:
                paragraph.style = "Code"
            except KeyError:
                pass
            paragraph.add_run(source)
