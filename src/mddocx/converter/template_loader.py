"""Load Word templates for conversion while preserving headers/footers."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn

_ALLOWED_SUFFIXES = {".docx", ".dotx"}


def resolve_template_path(template_path: str | Path) -> Path:
    path = Path(template_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"模板文件不存在: {path}")
    if path.suffix.lower() not in _ALLOWED_SUFFIXES:
        raise ValueError(f"不支持的模板格式（需 .docx 或 .dotx）: {path.suffix}")
    return path


def load_document_from_template(template_path: str | Path) -> DocxDocument:
    """Open a template and remove body content, keeping headers/footers/styles."""
    path = resolve_template_path(template_path)
    doc = Document(str(path))
    _clear_document_body(doc)
    return doc


def _clear_document_body(doc: DocxDocument) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)

    for table in list(doc.tables):
        tbl = table._element
        if tbl.getparent() is not None:
            tbl.getparent().remove(tbl)
    for paragraph in list(doc.paragraphs):
        p = paragraph._element
        if p.getparent() is not None:
            p.getparent().remove(p)
