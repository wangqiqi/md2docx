"""Cover page rendering from front matter metadata."""

from __future__ import annotations

from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from .front_matter import FrontMatter


def _style_or_none(document: DocxDocument, name: str):
    try:
        return document.styles[name]
    except KeyError:
        return None


def _add_centered(document: DocxDocument, text: str, style_name: str, *, space_before_pt: int = 0) -> Paragraph:
    paragraph = document.add_paragraph()
    style = _style_or_none(document, style_name)
    if style is not None:
        paragraph.style = style
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if space_before_pt:
        paragraph.paragraph_format.space_before = Pt(space_before_pt)
    paragraph.add_run(text)
    return paragraph


def render_cover_page(document: DocxDocument, meta: FrontMatter, title: str) -> None:
    """Insert a simple cover (title, optional author/date) and page break before body."""
    _add_centered(document, title, "Title", space_before_pt=120)
    subtitle_parts = []
    if meta.author:
        subtitle_parts.append(meta.author)
    if meta.date:
        subtitle_parts.append(meta.date)
    if subtitle_parts:
        _add_centered(document, " · ".join(subtitle_parts), "Subtitle", space_before_pt=24)

    break_para = document.add_paragraph()
    run = break_para.add_run()
    run.add_break(WD_BREAK.PAGE)
