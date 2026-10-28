"""WebUI convert with optional Word template (SPRINT-STYLE-03)."""

from __future__ import annotations

import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from mddocx.webui.app import app
from mddocx.webui.rate_limit import reset_rate_limits

TEMPLATE = (
    Path(__file__).resolve().parents[4] / "tests" / "samples" / "templates" / "minimal_company.docx"
)


@pytest.fixture
def client():
    app.config["TESTING"] = True
    app.config["WTF_CSRF_ENABLED"] = False
    reset_rate_limits()
    with app.test_client() as test_client:
        yield test_client
    reset_rate_limits()


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="sample template missing")
def test_convert_with_template_upload(client):
    with TEMPLATE.open("rb") as handle:
        response = client.post(
            "/convert",
            data={
                "markdown": "# WebUI Template\n\nBody line.",
                "template": (handle, "company.docx"),
            },
            content_type="multipart/form-data",
        )
    assert response.status_code == 200
    doc = Document(BytesIO(response.data))
    header = doc.sections[0].header.paragraphs[0].text
    assert "MD2DOCX-TEMPLATE-HEADER" in header
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "WebUI Template" in body


def test_convert_rejects_invalid_template_extension(client):
    response = client.post(
        "/convert",
        data={
            "markdown": "# X",
            "template": (BytesIO(b"not-a-docx"), "evil.pdf"),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"E_FILE_TYPE_INVALID" in response.data


def test_convert_without_template_unchanged(client):
    response = client.post("/convert", data={"markdown": "# Plain\n\nok."})
    assert response.status_code == 200
    doc = Document(BytesIO(response.data))
    assert "Plain" in doc.paragraphs[0].text
