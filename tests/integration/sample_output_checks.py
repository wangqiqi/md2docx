"""
tests/samples 导出 DOCX 程序化验收规则（对照源 .md 的期望，非 golden 二进制对比）。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Callable, List, Optional
from zipfile import ZipFile

from docx import Document

try:
    from PIL import Image
except ImportError:
    Image = None  # type: ignore


@dataclass
class SampleExpectation:
    """单一样例 docx 的验收规则。"""

    min_size: int = 8_000
    min_paragraphs: int = 1
    min_tables: int = 0
    min_media: int = 0
    text_contains: List[str] = field(default_factory=list)
    text_contains_any: List[str] = field(default_factory=list)
    custom: Optional[Callable[[Document, str, List[bytes]], List[str]]] = None


def first_h1(md_path: Path) -> Optional[str]:
    with open(md_path, "r", encoding="utf-8") as f:
        for line in f:
            if line.startswith("# "):
                return line[2:].strip()
    return None


def paragraph_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def table_cell_text(doc: Document) -> str:
    parts: List[str] = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def full_text(doc: Document) -> str:
    return paragraph_text(doc) + "\n" + table_cell_text(doc)


def extract_media_bytes(docx_path: Path) -> List[bytes]:
    blobs: List[bytes] = []
    with ZipFile(docx_path) as zf:
        for name in zf.namelist():
            if name.startswith("word/media/"):
                blobs.append(zf.read(name))
    return blobs


def validate_png_or_jpeg(blob: bytes) -> bool:
    if len(blob) < 8:
        return False
    if blob.startswith(b"\x89PNG\r\n\x1a\n") or blob.startswith(b"\xff\xd8\xff"):
        if Image is None:
            return True
        try:
            Image.open(BytesIO(blob)).verify()
            return True
        except Exception:
            return False
    return len(blob) >= 100


# 键为相对 tests/samples/ 的路径（POSIX）
EXPECTATIONS: dict[str, SampleExpectation] = {
    "basic/headings.md": SampleExpectation(
        min_paragraphs=5,
        text_contains=["这是一级标题", "二级标题"],
    ),
    "basic/text_styles.md": SampleExpectation(
        text_contains=["粗体和斜体", "删除线"],
    ),
    "basic/lists.md": SampleExpectation(
        min_paragraphs=10,
        text_contains=["无序列表", "有序列表"],
    ),
    "basic/blockquotes.md": SampleExpectation(
        text_contains=["引用块示例", "外层引用", "内层引用"],
    ),
    "basic/code.md": SampleExpectation(
        min_paragraphs=5,
        text_contains_any=["代码", "python", "def "],
    ),
    "basic/tables.md": SampleExpectation(
        min_tables=1,
        text_contains=["表格测试"],
    ),
    "basic/links.md": SampleExpectation(
        text_contains=["链接示例", "示例链接"],
    ),
    "basic/image.md": SampleExpectation(
        min_media=1,
        text_contains=["图片测试", "本地图片", "内联图片"],
    ),
    "basic/html.md": SampleExpectation(
        min_paragraphs=10,
        text_contains=["HTML", "div"],
    ),
    "basic/hr_and_tasks.md": SampleExpectation(
        min_paragraphs=20,
        text_contains=["分隔线", "任务列表"],
    ),
    "advanced/math.md": SampleExpectation(
        min_media=4,
        min_tables=1,
        text_contains=["数学公式", "(5)", "由上式 (5)"],
    ),
    "advanced/flowcharts.md": SampleExpectation(
        min_media=5,
        text_contains=["流程图", "时序图", "状态图", "甘特图", "饼图"],
    ),
    "advanced/tables.md": SampleExpectation(
        min_tables=3,
        text_contains=["表格测试"],
    ),
    "large/chunked.md": SampleExpectation(
        min_tables=1,
        text_contains=["大文档分块", "第一节", "第三节"],
    ),
    "test.md": SampleExpectation(
        min_paragraphs=50,
        min_tables=2,
        text_contains=[
            "Markdown to DOCX",
            "粗体文本",
            "任务列表",
            "代码块测试",
        ],
    ),
}


def check_docx_against_expectation(
    docx_path: Path,
    md_path: Path,
    exp: SampleExpectation,
) -> List[str]:
    """返回失败原因列表；空列表表示通过。"""
    failures: List[str] = []

    if not docx_path.is_file():
        return [f"docx 不存在: {docx_path}"]

    size = docx_path.stat().st_size
    if size < exp.min_size:
        failures.append(f"体积过小: {size} < {exp.min_size}")

    try:
        with ZipFile(docx_path) as zf:
            if "word/document.xml" not in zf.namelist():
                failures.append("非合法 docx（缺 word/document.xml）")
    except Exception as exc:
        failures.append(f"无法作为 ZIP 打开: {exc}")
        return failures

    doc = Document(str(docx_path))
    text = full_text(doc)
    media_blobs = extract_media_bytes(docx_path)

    h1 = first_h1(md_path)
    if h1 and h1 not in text:
        failures.append(f"缺少一级标题文本: {h1!r}")

    if len(doc.paragraphs) < exp.min_paragraphs:
        failures.append(
            f"段落过少: {len(doc.paragraphs)} < {exp.min_paragraphs}"
        )

    if len(doc.tables) < exp.min_tables:
        failures.append(f"表格过少: {len(doc.tables)} < {exp.min_tables}")

    if len(media_blobs) < exp.min_media:
        failures.append(f"嵌入媒体过少: {len(media_blobs)} < {exp.min_media}")

    for needle in exp.text_contains:
        if needle not in text:
            failures.append(f"正文缺少期望文本: {needle!r}")

    if exp.text_contains_any:
        if not any(n in text for n in exp.text_contains_any):
            failures.append(
                f"正文未命中任一期望: {exp.text_contains_any}"
            )

    for i, blob in enumerate(media_blobs, 1):
        if not validate_png_or_jpeg(blob):
            failures.append(f"word/media 第 {i} 项不是有效图片 ({len(blob)} B)")

    if exp.custom:
        failures.extend(exp.custom(doc, text, media_blobs))

    return failures


def list_sample_md_files(samples_root: Path) -> List[Path]:
    return sorted(
        p
        for p in samples_root.rglob("*.md")
        if p.is_file() and p.parent.name != "output"
    )


def rel_sample_key(md_path: Path, samples_root: Path) -> str:
    return md_path.relative_to(samples_root).as_posix()
