"""
BaseConverter token 路由边界测试
"""

from unittest.mock import MagicMock, patch

import pytest
from docx import Document

from mddocx.converter.base import BaseConverter

FAKE_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f"
    b"\x00\x00\x01\x01\x00\x05\x18\xd8d\x00\x00\x00\x00IEND\xaeB`\x82"
)


class TestBaseConverterRouting:
    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.mermaid.requests.get")
    def test_fence_mermaid_routes_to_mermaid(self, mock_get, mock_add_picture):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        md = "```mermaid\ngraph TD\n  A --> B\n```"
        doc = BaseConverter().convert(md)
        assert mock_get.called

    def test_fence_python_routes_to_code(self):
        md = "```python\nprint('hi')\n```"
        doc = BaseConverter().convert(md)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "print('hi')" in text

    def test_fence_without_lang_routes_to_code(self):
        md = "```\nplain code\n```"
        doc = BaseConverter().convert(md)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "plain code" in text

    def test_html_block_routes_to_html_converter(self):
        md = "<div>块级 HTML</div>\n\n正文"
        doc = BaseConverter().convert(md)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "块级 HTML" in text or "正文" in text

    def test_html_inline_in_paragraph(self):
        md = "段落含 <strong>粗体</strong> 文本"
        doc = BaseConverter().convert(md)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "粗体" in text
