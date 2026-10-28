"""Tests for Word template-based conversion."""

from pathlib import Path

import pytest
from docx import Document

from mddocx.converter.base import BaseConverter, ConvertError
from mddocx.converter.converter_pool import reset_converter_pool

TEMPLATE = Path(__file__).resolve().parents[1] / "samples" / "templates" / "minimal_company.docx"
HEADER_MARKER = "MD2DOCX-TEMPLATE-HEADER"


@pytest.fixture(autouse=True)
def _reset_pool():
    reset_converter_pool()
    yield
    reset_converter_pool()


def test_template_preserves_header_in_output(tmp_path):
    assert TEMPLATE.is_file()
    md = tmp_path / "body.md"
    md.write_text("# Title\n\nBody from markdown.\n", encoding="utf-8")
    out = tmp_path / "out.docx"

    doc = BaseConverter().convert_file(md, template_path=str(TEMPLATE))
    doc.save(out)

    saved = Document(str(out))
    header_text = saved.sections[0].header.paragraphs[0].text
    assert HEADER_MARKER in header_text
    body = "\n".join(p.text for p in saved.paragraphs)
    assert "Title" in body
    assert "Body from markdown" in body
    assert HEADER_MARKER not in body


def test_template_missing_raises():
    with pytest.raises(ConvertError, match="模板文件不存在"):
        BaseConverter().convert("# x", template_path="/nonexistent/template.docx")


def test_template_with_chunked_raises():
    with pytest.raises(ConvertError, match="chunked"):
        BaseConverter().convert("# x", template_path=str(TEMPLATE), chunked=True)
