"""
引用块转换器测试
"""

import pytest
from docx import Document

from mddocx.converter.base import BaseConverter
from mddocx.converter.elements.blockquote import BlockquoteConverter


class TestBlockquoteConverter:
    """引用块转换器测试类"""

    @pytest.fixture
    def converter(self):
        """创建带链接转换器的引用块转换器"""
        base = BaseConverter()
        conv = BlockquoteConverter(base)
        conv.set_document(Document())
        return conv

    def test_basic_blockquote(self, converter):
        """测试基本引用块"""
        doc = BaseConverter().convert("> Hello quote")
        assert len(doc.paragraphs) >= 1
        assert "Hello quote" in doc.paragraphs[0].text

    def test_blockquote_with_link(self, converter):
        """测试引用块内链接"""
        doc = BaseConverter().convert("> See [example](https://example.com) here")
        para = doc.paragraphs[0]
        hyperlinks = para._element.xpath(".//w:hyperlink")
        assert len(hyperlinks) >= 1

    def test_blockquote_with_strikethrough(self, converter):
        """测试引用块内删除线"""
        doc = BaseConverter().convert("> ~~removed~~ text")
        runs = doc.paragraphs[0].runs
        assert any(run.font.strike for run in runs if run.text)

    def test_blockquote_with_inline_code(self, converter):
        """测试引用块内行内代码"""
        doc = BaseConverter().convert("> use `foo()` here")
        runs = doc.paragraphs[0].runs
        assert any("foo()" in run.text for run in runs)
        assert any(run.font.name == "Consolas" for run in runs)

    def test_blockquote_mixed_inline(self, converter):
        """测试引用块混合内联元素"""
        md = "> **bold** [link](https://a.com) `code` ~~strike~~"
        doc = BaseConverter().convert(md)
        text = doc.paragraphs[0].text
        assert "bold" in text
        assert "code" in text
        assert "strike" in text
        hyperlinks = doc.paragraphs[0]._element.xpath(".//w:hyperlink")
        assert len(hyperlinks) >= 1
