"""
数学公式转换器单元测试
"""

from unittest.mock import MagicMock, patch

import pytest
from docx import Document

from mddocx.converter.base import BaseConverter
from mddocx.converter.elements.math import MathConverter, build_codecogs_url
from mddocx.converter.equation_labels import EquationRegistry, strip_label, substitute_refs
from mddocx.converter.security import is_allowed_codecogs_url

FAKE_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f"
    b"\x00\x00\x01\x01\x00\x05\x18\xd8d\x00\x00\x00\x00IEND\xaeB`\x82"
)


class TestMathHelpers:
    def test_build_codecogs_url(self):
        url = build_codecogs_url(r"\frac{1}{2}")
        assert url.startswith("https://latex.codecogs.com/png.latex?")
        assert "frac" in url

    def test_build_codecogs_inline_prefix(self):
        url = build_codecogs_url("E=mc^2", inline=True)
        assert "%5Cinline" in url or "\\inline" in url


class TestMathSecurity:
    def test_allowed_codecogs_url(self):
        url = build_codecogs_url(r"\alpha")
        assert is_allowed_codecogs_url(url) is True

    def test_rejects_arbitrary_host(self):
        assert is_allowed_codecogs_url("https://evil.com/png.latex?x") is False

    def test_rejects_http(self):
        url = "http://latex.codecogs.com/png.latex?x"
        assert is_allowed_codecogs_url(url) is False


class TestMathConverter:
    @pytest.fixture
    def converter(self):
        base = BaseConverter()
        conv = base.converters["math"]
        conv.set_document(Document())
        return conv

    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.math.requests.get")
    def test_render_block(self, mock_get, mock_add_picture, converter):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        token = MagicMock()
        token.content = r"\sum_{i=1}^n i"
        converter.convert(token)

        assert mock_get.called
        mock_add_picture.assert_called_once()

    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.math.requests.get")
    def test_render_inline(self, mock_get, mock_add_picture, converter):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        para = converter.document.add_paragraph()
        token = MagicMock()
        token.content = "E=mc^2"
        converter.convert_in_paragraph(para, token)

        assert mock_get.called
        mock_add_picture.assert_called_once()

    @patch("mddocx.converter.elements.math.requests.get")
    def test_render_failure_fallback(self, mock_get, converter):
        mock_get.side_effect = Exception("network error")

        token = MagicMock()
        token.content = r"\frac{1}{2}"
        converter.convert(token)

        text = "\n".join(p.text for p in converter.document.paragraphs)
        assert "渲染失败" in text
        assert "frac" in text


class TestMathRouting:
    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.math.requests.get")
    def test_base_converter_inline_math(self, mock_get, mock_add_picture):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        doc = BaseConverter().convert("Energy $E=mc^2$ formula.")
        assert mock_get.called
        assert len(doc.paragraphs) >= 1

    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.math.requests.get")
    def test_base_converter_block_math(self, mock_get, mock_add_picture):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        doc = BaseConverter().convert(r"$$\int_0^1 x\\,dx$$\n")
        assert mock_get.called
        assert len(doc.paragraphs) >= 1


class TestEquationLabels:
    def test_strip_label_from_latex(self):
        cleaned, label = strip_label(r"E=mc^2 \label{eq:emc}")
        assert label == "eq:emc"
        assert r"\label" not in cleaned
        assert "E=mc^2" in cleaned

    def test_strip_label_without_label(self):
        cleaned, label = strip_label(r"\frac{1}{2}")
        assert label is None
        assert cleaned == r"\frac{1}{2}"

    def test_registry_register_and_resolve(self):
        reg = EquationRegistry()
        reg.register("eq:a", 1)
        assert reg.resolve("eq:a") == 1
        assert reg.resolve_ref_text("eq:a") == "(1)"
        assert reg.resolve_ref_text("missing") == "(?)"


class TestEquationNumbering:
    @pytest.fixture
    def converter(self):
        base = BaseConverter()
        conv = base.converters["math"]
        conv.set_document(Document())
        return conv

    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.math.requests.get")
    def test_block_equation_number_caption(self, mock_get, mock_add_picture, converter):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        for content in (r"\alpha", r"\beta"):
            token = MagicMock()
            token.content = content
            converter.convert(token)

        text = "\n".join(p.text for p in converter.document.paragraphs)
        assert "(1)" in text
        assert "(2)" in text

    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.math.requests.get")
    def test_label_stripped_from_codecogs_request(self, mock_get, mock_add_picture, converter):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        token = MagicMock()
        token.content = r"E=mc^2 \label{eq:emc}"
        converter.convert(token)

        url = mock_get.call_args[0][0]
        assert "label" not in url.lower()


class TestEquationRef:
    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.math.requests.get")
    def test_ref_resolves_after_labeled_block(self, mock_get, mock_add_picture):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        md = r"""
$$
E=mc^2 \label{eq:emc}
$$

见式 \ref{eq:emc} 所示。
"""
        doc = BaseConverter().convert(md)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "(1)" in text
        assert "见式 (1) 所示" in text

    def test_ref_unknown_shows_placeholder(self):
        doc = BaseConverter().convert("引用 \\ref{unknown} 结束。")
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "(?)" in text

    def test_substitute_refs_helper(self):
        reg = EquationRegistry()
        reg.register("eq:x", 3)
        assert substitute_refs("式 \\ref{eq:x} 与 \\ref{bad}", reg) == "式 (3) 与 (?)"
