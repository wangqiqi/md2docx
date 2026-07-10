"""
Mermaid 转换器单元测试
"""

from unittest.mock import MagicMock, patch

import pytest
from docx import Document

from mddocx.converter.base import BaseConverter
from mddocx.converter.elements.mermaid import (
    MermaidConverter,
    _is_valid_image_payload,
    build_mermaid_ink_url,
    is_supported_mermaid_diagram,
    mermaid_diagram_kind,
)
from mddocx.converter.security import MAX_IMAGE_BYTES, is_allowed_mermaid_ink_url

SAMPLE_GRAPH = """graph TD
    A[开始] --> B[结束]
"""

SAMPLE_SEQUENCE = """sequenceDiagram
    A->>B: hi
"""

SAMPLE_GANTT = """gantt
    title Plan
    section S1
    Task1 :2024-01-01, 7d
"""

SAMPLE_STATE = """stateDiagram-v2
    [*] --> A
    A --> [*]
"""

SAMPLE_CLASS = """classDiagram
    Animal <|-- Dog
"""

SAMPLE_PIE = """pie title 分配
    "A" : 40
    "B" : 60
"""

# 最小合法 1x1 PNG
FAKE_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f"
    b"\x00\x00\x01\x01\x00\x05\x18\xd8d\x00\x00\x00\x00IEND\xaeB`\x82"
)


class TestMermaidHelpers:
    def test_supported_graph_td(self):
        assert is_supported_mermaid_diagram(SAMPLE_GRAPH) is True

    def test_supported_flowchart(self):
        assert is_supported_mermaid_diagram("flowchart LR\n  A --> B") is True

    def test_supported_sequence(self):
        assert is_supported_mermaid_diagram(SAMPLE_SEQUENCE) is True

    def test_supported_gantt(self):
        assert is_supported_mermaid_diagram(SAMPLE_GANTT) is True

    def test_supported_state_diagram(self):
        assert is_supported_mermaid_diagram(SAMPLE_STATE) is True

    def test_supported_class_diagram(self):
        assert is_supported_mermaid_diagram(SAMPLE_CLASS) is True

    def test_supported_pie(self):
        assert is_supported_mermaid_diagram(SAMPLE_PIE) is True

    def test_unsupported_journey(self):
        assert is_supported_mermaid_diagram("journey\n  title: X") is False

    def test_diagram_kind_labels(self):
        assert mermaid_diagram_kind(SAMPLE_GRAPH) == "流程图"
        assert mermaid_diagram_kind(SAMPLE_SEQUENCE) == "时序图"
        assert mermaid_diagram_kind(SAMPLE_GANTT) == "甘特图"
        assert mermaid_diagram_kind(SAMPLE_STATE) == "状态图"
        assert mermaid_diagram_kind(SAMPLE_CLASS) == "类图"
        assert mermaid_diagram_kind(SAMPLE_PIE) == "饼图"

    def test_diagram_kind_empty_or_unknown(self):
        assert mermaid_diagram_kind("") == "图表"
        assert mermaid_diagram_kind("   ") == "图表"
        assert mermaid_diagram_kind("journey\n  title: X") == "图表"

    def test_is_supported_empty_source(self):
        assert is_supported_mermaid_diagram("") is False
        assert is_supported_mermaid_diagram("   ") is False

    def test_is_valid_image_payload(self):
        assert _is_valid_image_payload(b"") is False
        assert _is_valid_image_payload(b"short") is False
        assert _is_valid_image_payload(FAKE_PNG) is True
        assert _is_valid_image_payload(b"\xff\xd8\xff" + b"x" * 100) is True
        assert _is_valid_image_payload(b"x" * 100) is True

    def test_build_url_uses_mermaid_ink(self):
        url = build_mermaid_ink_url(SAMPLE_GRAPH)
        assert url.startswith("https://mermaid.ink/img/")
        assert "type=png" in url


class TestMermaidSecurity:
    def test_allowed_mermaid_ink_url(self):
        url = build_mermaid_ink_url(SAMPLE_GRAPH)
        assert is_allowed_mermaid_ink_url(url) is True

    def test_rejects_arbitrary_host(self):
        assert is_allowed_mermaid_ink_url("https://evil.com/img/abc") is False

    def test_rejects_non_https(self):
        assert is_allowed_mermaid_ink_url("http://mermaid.ink/img/abc") is False


class TestMermaidConverter:
    @pytest.fixture
    def converter(self):
        base = BaseConverter()
        conv = base.converters["mermaid"]
        conv.set_document(Document())
        return conv

    def _make_token(self, content: str, lang: str = "mermaid"):
        token = MagicMock()
        token.content = content
        token.info = lang
        return token

    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.mermaid.requests.get")
    def test_render_graph_embeds_image(self, mock_get, mock_add_picture, converter):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        converter.convert(self._make_token(SAMPLE_GRAPH))

        assert mock_get.called
        mock_add_picture.assert_called_once()
        text = "\n".join(p.text for p in converter.document.paragraphs)
        assert "Mermaid 流程图" in text

    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.mermaid.requests.get")
    def test_render_sequence_embeds_image(self, mock_get, mock_add_picture, converter):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        converter.convert(self._make_token(SAMPLE_SEQUENCE))

        mock_add_picture.assert_called_once()
        text = "\n".join(p.text for p in converter.document.paragraphs)
        assert "Mermaid 时序图" in text

    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.mermaid.requests.get")
    def test_render_gantt_embeds_image(self, mock_get, mock_add_picture, converter):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        converter.convert(self._make_token(SAMPLE_GANTT))

        mock_add_picture.assert_called_once()
        text = "\n".join(p.text for p in converter.document.paragraphs)
        assert "Mermaid 甘特图" in text

    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.mermaid.requests.get")
    def test_render_state_embeds_image(self, mock_get, mock_add_picture, converter):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        converter.convert(self._make_token(SAMPLE_STATE))
        mock_add_picture.assert_called_once()
        text = "\n".join(p.text for p in converter.document.paragraphs)
        assert "Mermaid 状态图" in text

    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.mermaid.requests.get")
    def test_render_class_embeds_image(self, mock_get, mock_add_picture, converter):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        converter.convert(self._make_token(SAMPLE_CLASS))
        mock_add_picture.assert_called_once()
        text = "\n".join(p.text for p in converter.document.paragraphs)
        assert "Mermaid 类图" in text

    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.mermaid.requests.get")
    def test_render_pie_embeds_image(self, mock_get, mock_add_picture, converter):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        converter.convert(self._make_token(SAMPLE_PIE))
        mock_add_picture.assert_called_once()
        text = "\n".join(p.text for p in converter.document.paragraphs)
        assert "Mermaid 饼图" in text

    @patch("mddocx.converter.elements.mermaid.requests.get")
    def test_render_failure_fallback(self, mock_get, converter):
        mock_get.side_effect = Exception("network error")

        converter.convert(self._make_token(SAMPLE_GRAPH))

        text = "\n".join(p.text for p in converter.document.paragraphs)
        assert "渲染失败" in text
        assert "graph TD" in text

    def test_unsupported_type_fallback(self, converter):
        converter.convert(self._make_token("journey\n  title: Trip"))

        text = "\n".join(p.text for p in converter.document.paragraphs)
        assert "不支持" in text
        assert "journey" in text

    def test_document_not_set(self):
        conv = MermaidConverter()
        with pytest.raises(ValueError, match="Document not set"):
            conv.convert(self._make_token(SAMPLE_GRAPH))

    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.mermaid.requests.get")
    def test_embed_failure_falls_back_to_code(self, mock_get, mock_add_picture, converter):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp
        mock_add_picture.side_effect = RuntimeError("embed fail")

        converter.convert(self._make_token(SAMPLE_GRAPH))

        text = "\n".join(p.text for p in converter.document.paragraphs)
        assert "渲染失败" in text

    @patch("mddocx.converter.elements.mermaid.requests.get")
    def test_fetch_blocked_url(self, mock_get, converter):
        with patch(
            "mddocx.converter.elements.mermaid.is_allowed_mermaid_ink_url",
            return_value=False,
        ):
            converter.convert(self._make_token(SAMPLE_GRAPH))

        mock_get.assert_not_called()
        text = "\n".join(p.text for p in converter.document.paragraphs)
        assert "渲染失败" in text

    @patch("mddocx.converter.elements.mermaid.requests.get")
    def test_fetch_non_200(self, mock_get, converter):
        mock_resp = MagicMock()
        mock_resp.status_code = 500
        mock_get.return_value = mock_resp

        converter.convert(self._make_token(SAMPLE_GRAPH))

        assert "渲染失败" in "\n".join(p.text for p in converter.document.paragraphs)

    @patch("mddocx.converter.elements.mermaid.requests.get")
    def test_fetch_bad_content_type(self, mock_get, converter):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "text/html"}
        mock_get.return_value = mock_resp

        converter.convert(self._make_token(SAMPLE_GRAPH))

        assert "渲染失败" in "\n".join(p.text for p in converter.document.paragraphs)

    @patch("mddocx.converter.elements.mermaid.requests.get")
    def test_fetch_oversized_image(self, mock_get, converter):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        chunk = b"x" * 8192
        mock_resp.iter_content.return_value = [chunk] * (MAX_IMAGE_BYTES // 8192 + 2)
        mock_get.return_value = mock_resp

        converter.convert(self._make_token(SAMPLE_GRAPH))

        assert "渲染失败" in "\n".join(p.text for p in converter.document.paragraphs)

    def test_fallback_without_code_converter(self, converter):
        converter.base_converter = None
        converter.convert(self._make_token("journey\n  title: Trip"))

        text = "\n".join(p.text for p in converter.document.paragraphs)
        assert "不支持" in text
        assert "journey" in text

    def test_fallback_plain_run_when_code_style_missing(self, converter):
        """无 code converter 且 Code 样式不存在时仍写入源码。"""
        converter.base_converter = MagicMock()
        converter.base_converter.converters = {}
        token = self._make_token(SAMPLE_GRAPH)

        paragraph = converter.doc.add_paragraph()
        with patch.object(
            type(paragraph),
            "style",
            property(
                lambda self: (_ for _ in ()).throw(KeyError("Code")),
                lambda self, value: None,
            ),
        ), patch.object(converter.doc, "add_paragraph", return_value=paragraph):
            converter._fallback_as_code(token, SAMPLE_GRAPH, unsupported=False)

        assert SAMPLE_GRAPH in "\n".join(p.text for p in converter.document.paragraphs)


class TestMermaidRouting:
    @patch("docx.text.run.Run.add_picture")
    @patch("mddocx.converter.elements.mermaid.requests.get")
    def test_base_converter_routes_mermaid_fence(self, mock_get, mock_add_picture):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.headers = {"Content-Type": "image/png"}
        mock_resp.iter_content.return_value = [FAKE_PNG]
        mock_get.return_value = mock_resp

        md = f"```mermaid\n{SAMPLE_GRAPH}\n```"
        doc = BaseConverter().convert(md)
        assert mock_get.called
        assert len(doc.paragraphs) >= 1
